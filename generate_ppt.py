import sys
import subprocess
import os
import urllib.request

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError:
    print("python-pptx not found. Installing...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

# Download a demo picture for the presentation
photo_path = "demo_photo.jpg"
if not os.path.exists(photo_path):
    print("Downloading placeholder photo...")
    try:
        # A nice placeholder image of a mobile phone / app to represent the app
        req = urllib.request.Request(
            "https://images.unsplash.com/photo-1512941937669-90a1b58e7e9c?q=80&w=800&auto=format&fit=crop",
            data=None, 
            headers={'User-Agent': 'Mozilla/5.0'}
        )
        with urllib.request.urlopen(req) as response, open(photo_path, 'wb') as out_file:
            data = response.read()
            out_file.write(data)
    except Exception as e:
        print(f"Failed to download image: {e}")

prs = Presentation()

def set_slide_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_slide(title_text, bullet_points, bg_color=RGBColor(245, 245, 250)):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, bg_color)
    
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(20, 20, 100) # Dark Blue title
    title.text_frame.paragraphs[0].font.bold = True
    
    body = slide.placeholders[1]
    tf = body.text_frame
    
    # Set first bullet point
    tf.text = bullet_points[0]
    tf.paragraphs[0].font.color.rgb = RGBColor(40, 40, 40)
    tf.paragraphs[0].font.size = Pt(24)
    
    # Set remaining bullet points
    for pt in bullet_points[1:]:
        p = tf.add_paragraph()
        p.text = pt
        p.level = 0
        p.font.color.rgb = RGBColor(40, 40, 40)
        p.font.size = Pt(24)
    
    return slide

# Slide 1: Title
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
set_slide_background(slide, RGBColor(33, 150, 243)) # Vibrant Blue

title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "FindIt: Lost and Found Application"
title.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255) # White text
title.text_frame.paragraphs[0].font.bold = True

subtitle.text = "A Streamlined Lost and Found Management System\nConnecting people with their lost belongings."
for paragraph in subtitle.text_frame.paragraphs:
    paragraph.font.color.rgb = RGBColor(230, 240, 255)
    paragraph.font.size = Pt(22)

# Slide 2: Introduction & Overview
add_slide("Introduction & Overview", [
    "Objective: Provide a digital platform to easily report, search for, and claim lost and found items.",
    "Problem Solved: Traditional lost and found processes are disorganized and localized.",
    "Solution: A centralized mobile application accessible anywhere, bringing together users who lost items with those who found them.",
    "Core Platform: Android App powered by Firebase and Cloudinary."
], RGBColor(255, 245, 230)) # Warm light orange bg

# Slide 3: System Architecture & Technologies
slide3 = add_slide("System Architecture & Technologies", [
    "Frontend: Android (Java/Kotlin) providing an intuitive native user interface.",
    "Backend/Database: Firebase Firestore for real-time, scalable data synchronization.",
    "Image Hosting: Cloudinary integration for robust handling of item images and proof of ownership.",
    "Architecture: Asynchronous data layer for seamless non-blocking operations."
])

# Insert photo on this slide if it exists
if os.path.exists(photo_path):
    # Place photo towards the bottom right
    left = Inches(5.5)
    top = Inches(4.5)
    height = Inches(2.5)
    try:
        slide3.shapes.add_picture(photo_path, left, top, height=height)
    except Exception as e:
        print("Could not add picture:", e)

# Slide 4: Part 1 - User Profile & Item Management
add_slide("Part 1: User Profile & Item Management", [
    "Authentication: Secure user registration and login system.",
    "Profile: Manage personal details and contact information.",
    "Reporting: Users can report items they have lost or items they have found.",
    "Details: Upload images (via Cloudinary), add descriptions, location details, and categorization.",
    "Tracking: Users can track the status of the items they have reported."
], RGBColor(240, 255, 240)) # Mint green

# Slide 5: Part 2 - Item Viewing and Search
add_slide("Part 2: Item Viewing and Search", [
    "Home Feed: Centralized feed displaying all recently lost and found items.",
    "Detailed View: Click on any item to view full descriptions, high-quality images, and context.",
    "Search Capabilities: Search by keywords (e.g., 'iPhone', 'Keys').",
    "Filtering: Filter listings by category, date, and status (Lost vs. Found)."
])

# Slide 6: Part 3 - Claims and Notifications
add_slide("Part 3: Claims and Notifications", [
    "Initiation: Users who recognize their lost item can initiate a claim directly from the item's page.",
    "Proof of Ownership: Users upload verification images when submitting a claim.",
    "Alerts: Users receive updates on their claim status (Pending, Verified, Rejected).",
    "Communication: System bridges the gap between the claimant and the administration securely."
], RGBColor(255, 235, 245)) # Soft pink

# Slide 7: Part 4 - Admin and Reports
add_slide("Part 4: Admin and Reports", [
    "Admin Dashboard: Admins have a dedicated panel to oversee all platform activity.",
    "Claim Verification: Admins review submitted claims and the attached proof of ownership images.",
    "'Mark Recovered': Admins can finalize the process by marking an item as 'RECOVERED'.",
    "System Reports: Generate insights on app usage, total items recovered, and active user metrics."
])

# Slide 8: Future Enhancements & Roadmap
add_slide("Future Enhancements & Roadmap", [
    "Implement push notifications for real-time match alerts.",
    "Integrate location-based matching using Maps API (e.g., Google Maps).",
    "Add an in-app messaging system for direct user-to-user communication after a claim is verified."
], RGBColor(230, 240, 255)) # Soft blue

# Slide 9: Q&A
slide9 = add_slide("Q&A", [
    "Thank You!",
    "Questions and feedback are welcome."
], RGBColor(255, 200, 100)) # Vibrant orange/yellow

# Add picture to last slide as well
if os.path.exists(photo_path):
    left = Inches(1)
    top = Inches(4.5)
    height = Inches(2.5)
    try:
        slide9.shapes.add_picture(photo_path, left, top, height=height)
    except Exception as e:
        print("Could not add picture:", e)

output_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), "FindIt_Presentation_Colorful.pptx")
prs.save(output_file)
print(f"Colorful presentation generated successfully at: {output_file}")
